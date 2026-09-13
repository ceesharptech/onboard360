"""Text extraction module for RAG service.

Supports PDF (via PyMuPDF/fitz) and Word (.docx via python-docx).
Explicitly detects and flags image-only/scanned documents that have no extractable text.
"""

from pathlib import Path
import pymupdf as fitz
import docx


class ExtractionError(Exception):
    """Raised when text extraction fails or when a document contains no extractable text layer."""
    pass


def extract_text_from_pdf(file_path: str | Path) -> str:
    """Extract text from a PDF file using PyMuPDF.
    
    Raises ExtractionError if the document is empty or scanned (no text layer).
    """
    path = Path(file_path)
    if not path.exists():
        raise ExtractionError(f"File not found: {path}")

    try:
        doc = fitz.open(str(path))
    except Exception as exc:
        raise ExtractionError(f"Failed to open PDF: {str(exc)}") from exc

    page_texts: list[str] = []
    try:
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text")
            if text and text.strip():
                page_texts.append(text.strip())
    finally:
        doc.close()

    full_text = "\n\n".join(page_texts).strip()

    # Detect image-only or empty PDFs
    if not full_text:
        raise ExtractionError(
            "No extractable text found in document. Scanned or image-only PDFs are not supported."
        )

    return full_text


def extract_text_from_docx(file_path: str | Path) -> str:
    """Extract text from a Word (.docx) file including paragraphs and tables."""
    path = Path(file_path)
    if not path.exists():
        raise ExtractionError(f"File not found: {path}")

    try:
        doc = docx.Document(str(path))
    except Exception as exc:
        raise ExtractionError(f"Failed to open DOCX: {str(exc)}") from exc

    elements: list[str] = []

    for paragraph in doc.paragraphs:
        p_text = paragraph.text.strip()
        if p_text:
            elements.append(p_text)

    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                elements.append(" | ".join(row_cells))

    full_text = "\n\n".join(elements).strip()

    if not full_text:
        raise ExtractionError("No extractable text found in Word document.")

    return full_text


def extract_text(file_path: str | Path, filename: str | None = None) -> str:
    """Extract text from a document based on its file extension.
    
    Supports .pdf and .docx only.
    """
    path = Path(file_path)
    effective_name = filename or path.name
    suffix = Path(effective_name).suffix.lower()

    if suffix == ".pdf":
        return extract_text_from_pdf(path)
    elif suffix in (".docx", ".doc"):
        if suffix == ".doc":
            raise ExtractionError("Legacy .doc format is not supported. Please upload a .docx file.")
        return extract_text_from_docx(path)
    else:
        raise ExtractionError(
            f"Unsupported file format '{suffix}'. Only .pdf and .docx documents are supported."
        )
